"""
XLSB to XLSX Converter - Streamlit app.

Three modes:
  1. Single file    -> result shown in the UI, downloadable.
  2. Multiple files -> every result shown in the UI, each downloadable (+ ZIP).
  3. Folder         -> only .xlsb files are converted; results go into a new
                       "Outputs" folder inside the same folder.
                       * Running locally: type the folder path.
                       * Hosted (Streamlit Cloud): the server cannot see your
                         disk, so upload the folder and download a ZIP that
                         contains the "Outputs" folder.
"""

from __future__ import annotations

import io
import shutil
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
import streamlit as st

import converter as cv

APP_TITLE = "XLSB to XLSX Converter"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PREVIEW_ROWS = 50
PREVIEW_COLS = 30

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --------------------------------------------------------------------------
# Styling (dark theme itself lives in .streamlit/config.toml)
# --------------------------------------------------------------------------

st.markdown(
    """
    <style>
      #MainMenu, footer { visibility: hidden; }
      .block-container { padding-top: 2.2rem; padding-bottom: 3rem; max-width: 1100px; }

      .hero {
        background: linear-gradient(135deg, #102A43 0%, #0F3D3E 55%, #0B1220 100%);
        border: 1px solid #1F3A4D;
        border-radius: 16px;
        padding: 28px 32px;
        margin-bottom: 1.4rem;
      }
      .hero h1 { margin: 0 0 6px 0; font-size: 2rem; letter-spacing: -0.02em; color: #F1F5F9; }
      .hero p  { margin: 0; color: #9FB3C8; font-size: 1rem; line-height: 1.5; }
      .chip {
        display: inline-block; padding: 2px 12px; border-radius: 999px;
        font-size: .75rem; font-weight: 600; letter-spacing: .02em;
        border: 1px solid #2B6F5A; color: #34D399;
        background: rgba(52, 211, 153, .08); margin: 0 8px 12px 0;
      }
      .chip.muted { border-color: #33455A; color: #9FB3C8; background: rgba(159,179,200,.06); }

      .engine-line { font-size: .88rem; margin: 2px 0; }
      .dot { display:inline-block; width:8px; height:8px; border-radius:50%; margin-right:8px; }
      .dot.on  { background:#34D399; }
      .dot.off { background:#475569; }

      .section-title { font-size: 1.05rem; font-weight: 600; margin: .2rem 0 .6rem 0; }
      div[data-baseweb="tab-list"] { gap: 6px; }
    </style>
    """,
    unsafe_allow_html=True,
)


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def fmt_size(num_bytes: int) -> str:
    if num_bytes >= 1024 * 1024:
        return f"{num_bytes / (1024 * 1024):.2f} MB"
    return f"{max(num_bytes, 0) / 1024:.1f} KB"


def new_workdir(key: str) -> Path:
    """Fresh per-session scratch folder for a mode; removes the previous one."""
    old = st.session_state.get(f"{key}_dir")
    if old:
        shutil.rmtree(old, ignore_errors=True)
    path = Path(tempfile.mkdtemp(prefix=f"x2x_{key}_"))
    st.session_state[f"{key}_dir"] = str(path)
    return path


def run_jobs(jobs, engine_choice: str, timeout: int) -> list[dict]:
    """jobs = [(label, src_path, dst_path, relative_output_name)]"""
    records: list[dict] = []
    total = len(jobs)
    bar = st.progress(0.0, text="Starting...")
    for index, (label, src, dst, rel) in enumerate(jobs, start=1):
        bar.progress((index - 1) / total, text=f"Converting {index} of {total}: {label}")
        result = cv.convert_file(src, dst, engine_choice, timeout)
        records.append(
            {
                "file": label,
                "ok": result.ok,
                "engine": result.engine,
                "size": result.size_bytes,
                "seconds": result.seconds,
                "output": result.output_path,
                "rel": rel,
                "error": result.error,
            }
        )
    bar.progress(1.0, text="Finished")
    bar.empty()
    return records


def make_zip(records: list[dict], prefix: str = "") -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for rec in records:
            if rec["ok"]:
                archive.write(rec["output"], arcname=f"{prefix}{rec['rel']}")
    return buffer.getvalue()


def render_summary(records: list[dict], ignored: int = 0) -> None:
    ok = [r for r in records if r["ok"]]
    failed = [r for r in records if not r["ok"]]

    cols = st.columns(4 if ignored else 3)
    cols[0].metric("Converted", len(ok))
    cols[1].metric("Failed", len(failed))
    cols[2].metric("Total time", f"{sum(r['seconds'] for r in records):.1f} s")
    if ignored:
        cols[3].metric("Ignored (not .xlsb)", ignored)

    if any(r["ok"] and r["engine"] == cv.ENGINE_PYXLSB for r in records):
        st.warning(
            "Some files were converted with the **pyxlsb** fallback, which copies "
            "cell **values only** (no formulas, formatting or charts)."
        )

    table = pd.DataFrame(
        [
            {
                "File": r["file"],
                "Status": "✅ Converted" if r["ok"] else "❌ Failed",
                "Engine": cv.ENGINE_LABELS.get(r["engine"], r["engine"]) if r["ok"] else "",
                "Output size": fmt_size(r["size"]) if r["ok"] else "",
                "Time (s)": r["seconds"],
                "Details": r["error"],
            }
            for r in records
        ]
    )
    st.dataframe(table, hide_index=True)

    if failed:
        st.error(
            "Some files could not be converted. See the Details column. "
            "Try again, or choose another engine in the sidebar."
        )


def preview_workbook(path: str, key: str) -> None:
    """Small read-only preview of a converted workbook."""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = list(wb.sheetnames)
            sheet = st.selectbox("Sheet", sheets, key=f"{key}_sheet") if len(sheets) > 1 else sheets[0]
            ws = wb[sheet]
            rows = list(
                ws.iter_rows(
                    min_row=1,
                    max_row=PREVIEW_ROWS,
                    max_col=PREVIEW_COLS,
                    values_only=True,
                )
            )
        finally:
            wb.close()

        if not rows:
            st.caption("This sheet is empty.")
            return
        frame = pd.DataFrame(rows).fillna("").astype(str)
        st.caption(
            f"Showing up to the first {PREVIEW_ROWS} rows and {PREVIEW_COLS} columns "
            f"of '{sheet}' ({len(sheets)} sheet{'s' if len(sheets) != 1 else ''} in total)."
        )
        st.dataframe(frame, hide_index=True)
    except Exception as exc:  # noqa: BLE001
        st.caption(f"Preview unavailable: {exc}")


def download_list(records: list[dict], key: str) -> None:
    """One download button per converted file."""
    for i, rec in enumerate(r for r in records if r["ok"]):
        left, right = st.columns([5, 2], vertical_alignment="center")
        left.markdown(f"**{rec['rel']}** · {fmt_size(rec['size'])}")
        with open(rec["output"], "rb") as handle:
            right.download_button(
                "Download",
                data=handle.read(),
                file_name=Path(rec["rel"]).name,
                mime=XLSX_MIME,
                key=f"{key}_dl_{i}",
            )


# --------------------------------------------------------------------------
# Sidebar
# --------------------------------------------------------------------------

status = cv.engine_status()
available = [e for e in cv.ENGINE_ORDER if status[e][0]]

with st.sidebar:
    st.markdown("### ⚙️ Settings")

    if not available:
        st.error("No conversion engine is available on this machine.")
        st.stop()

    engine_choice = st.selectbox(
        "Conversion engine",
        options=["auto"] + available,
        format_func=lambda k: "Auto (recommended)" if k == "auto" else cv.ENGINE_LABELS[k],
        help="Auto picks the best available engine and falls back to the next one if a file fails.",
    )

    timeout_seconds = st.number_input(
        "Timeout per file (seconds)",
        min_value=30,
        max_value=3600,
        value=300,
        step=30,
        help="Enforced for LibreOffice. Microsoft Excel (COM) cannot be interrupted safely.",
    )

    st.divider()
    st.markdown("**Engines on this machine**")
    for key in cv.ENGINE_ORDER:
        ok, detail = status[key]
        st.markdown(
            f"<div class='engine-line'><span class='dot {'on' if ok else 'off'}'></span>"
            f"{cv.ENGINE_LABELS[key]} <span style='color:#64748B'>· {detail}</span></div>",
            unsafe_allow_html=True,
        )

    active_chain = cv.resolve_engine_chain(engine_choice)
    if active_chain and active_chain[0] == cv.ENGINE_PYXLSB:
        st.warning("Only the values-only fallback is available. Formulas and formatting will not be kept.")

    st.divider()
    st.caption(
        "XLSX cannot hold VBA macros. Always spot-check critical workbooks "
        "after conversion."
    )

active_label = cv.ENGINE_LABELS[active_chain[0]] if active_chain else "None"
hosted = cv.is_hosted_cloud()

# --------------------------------------------------------------------------
# Header
# --------------------------------------------------------------------------

st.markdown(
    f"""
    <div class="hero">
      <span class="chip">XLSB → XLSX</span>
      <span class="chip muted">Engine: {active_label}</span>
      <h1>{APP_TITLE}</h1>
      <p>Convert Excel Binary Workbooks to standard .xlsx. Convert one file,
      several files at once, or a whole folder.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

tab_single, tab_multi, tab_folder = st.tabs(
    ["📄  Single file", "🗂️  Multiple files", "📁  Folder"]
)

# --------------------------------------------------------------------------
# Mode 1: single file
# --------------------------------------------------------------------------

with tab_single:
    st.markdown("<div class='section-title'>Convert one workbook</div>", unsafe_allow_html=True)

    upload = st.file_uploader(
        "Drop an .xlsb file here",
        type=["xlsb"],
        accept_multiple_files=False,
        key="up_single",
    )

    if upload is None:
        st.session_state.pop("res_single", None)
    else:
        st.caption(f"{upload.name} · {fmt_size(upload.size)}")

        if st.button("Convert to XLSX", type="primary", key="go_single"):
            work = new_workdir("single")
            try:
                name = Path(cv.safe_relative_path(upload.name).name)
                src = work / name.name
                src.write_bytes(upload.getbuffer())
                dst = work / f"{name.stem}.xlsx"
                with st.spinner("Converting..."):
                    records = run_jobs(
                        [(name.name, src, dst, dst.name)], engine_choice, int(timeout_seconds)
                    )
                st.session_state["res_single"] = {"source": upload.name, "records": records}
            except Exception as exc:  # noqa: BLE001
                st.session_state.pop("res_single", None)
                st.error(f"Unable to process the upload: {exc}")

    state = st.session_state.get("res_single")
    if state and upload is not None and state["source"] == upload.name:
        rec = state["records"][0]
        st.divider()
        if rec["ok"]:
            with st.container(border=True):
                st.success("Conversion completed.")
                c1, c2, c3 = st.columns(3)
                c1.metric("Output", Path(rec["rel"]).name)
                c2.metric("Size", fmt_size(rec["size"]))
                c3.metric("Time", f"{rec['seconds']} s")
                with open(rec["output"], "rb") as handle:
                    st.download_button(
                        "⬇️  Download XLSX",
                        data=handle.read(),
                        file_name=Path(rec["rel"]).name,
                        mime=XLSX_MIME,
                        type="primary",
                        key="dl_single",
                    )
            if rec["engine"] == cv.ENGINE_PYXLSB:
                st.warning("Converted with the values-only fallback: formulas and formatting were not kept.")
            with st.expander("Preview converted data"):
                preview_workbook(rec["output"], "prev_single")
        else:
            st.error("Conversion failed.")
            st.code(rec["error"], language=None)

# --------------------------------------------------------------------------
# Mode 2: multiple files
# --------------------------------------------------------------------------

with tab_multi:
    st.markdown("<div class='section-title'>Convert several workbooks</div>", unsafe_allow_html=True)

    uploads = st.file_uploader(
        "Drop one or more .xlsb files here",
        type=["xlsb"],
        accept_multiple_files=True,
        key="up_multi",
    )

    if not uploads:
        st.session_state.pop("res_multi", None)
    else:
        total_size = sum(u.size for u in uploads)
        st.caption(f"{len(uploads)} file(s) selected · {fmt_size(total_size)}")

        if st.button("Convert all files", type="primary", key="go_multi"):
            work = new_workdir("multi")
            in_dir, out_dir = work / "in", work / "out"
            in_dir.mkdir()
            out_dir.mkdir()

            jobs, used = [], set()
            for index, item in enumerate(uploads):
                name = Path(item.name).name
                dst = cv.unique_output_path(out_dir, Path(name), used)
                src = in_dir / f"{index}_{name}"
                src.write_bytes(item.getbuffer())
                jobs.append((name, src, dst, dst.name))

            records = run_jobs(jobs, engine_choice, int(timeout_seconds))
            st.session_state["res_multi"] = {
                "count": len(uploads),
                "records": records,
                "zip": make_zip(records) if any(r["ok"] for r in records) else b"",
            }

    state = st.session_state.get("res_multi")
    if state and uploads and state["count"] == len(uploads):
        st.divider()
        render_summary(state["records"])

        if state["zip"]:
            st.download_button(
                "⬇️  Download all as ZIP",
                data=state["zip"],
                file_name="converted_xlsx_files.zip",
                mime="application/zip",
                type="primary",
                key="dl_multi_zip",
            )
            with st.expander("Download files individually", expanded=True):
                download_list(state["records"], "multi")

# --------------------------------------------------------------------------
# Mode 3: folder
# --------------------------------------------------------------------------

with tab_folder:
    st.markdown("<div class='section-title'>Convert every .xlsb in a folder</div>", unsafe_allow_html=True)
    st.caption("Only .xlsb files are converted. Files with any other extension are ignored.")

    source = st.radio(
        "Folder source",
        ["Folder path on this machine", "Upload a folder"],
        index=1 if hosted else 0,
        horizontal=True,
        key="folder_source",
    )

    # ---- 3a. server/local path ------------------------------------------
    if source == "Folder path on this machine":
        if hosted:
            st.warning(
                "This app is running on a remote server, so a path here refers to the "
                "**server's** disk, not your computer. Use **Upload a folder** to convert "
                "folders from your own computer."
            )
        else:
            st.caption(
                "Converted files are written to a new **Outputs** folder inside the folder "
                "you enter. Existing files with the same name in Outputs are replaced."
            )

        raw_path = st.text_input(
            "Folder path",
            placeholder=r"C:\Users\you\Documents\reports",
            key="folder_path",
        )

        folder = Path(raw_path.strip().strip('"')).expanduser() if raw_path.strip() else None
        found: list[Path] = []
        ignored = 0

        if folder is not None:
            if not folder.is_dir():
                st.error("That folder does not exist or is not accessible.")
            else:
                try:
                    found, ignored = cv.list_xlsb_files(folder)
                except OSError as exc:
                    st.error(f"Cannot read that folder: {exc}")

                if found:
                    st.success(
                        f"Found {len(found)} .xlsb file(s)"
                        + (f"; {ignored} other file(s) will be ignored." if ignored else ".")
                    )
                    with st.expander("Files to convert"):
                        st.dataframe(
                            pd.DataFrame(
                                {"File": [p.name for p in found], "Size": [fmt_size(p.stat().st_size) for p in found]}
                            ),
                            hide_index=True,
                        )
                elif folder.is_dir():
                    st.info("No .xlsb files were found directly inside this folder.")

        if found and st.button(
            f"Convert {len(found)} file(s) into Outputs", type="primary", key="go_folder_path"
        ):
            out_dir = folder / "Outputs"
            try:
                out_dir.mkdir(exist_ok=True)
            except OSError as exc:
                st.error(f"Cannot create the Outputs folder: {exc}")
            else:
                used: set[str] = set()
                jobs = []
                for src in found:
                    dst = cv.unique_output_path(out_dir, Path(src.name), used)
                    jobs.append((src.name, src, dst, dst.name))
                records = run_jobs(jobs, engine_choice, int(timeout_seconds))
                st.session_state["res_folder"] = {
                    "kind": "path",
                    "records": records,
                    "outputs": str(out_dir),
                    "ignored": ignored,
                }

        state = st.session_state.get("res_folder")
        if state and state["kind"] == "path":
            st.divider()
            ok_count = sum(r["ok"] for r in state["records"])
            if ok_count:
                st.success(f"{ok_count} file(s) saved to:")
                st.code(state["outputs"], language=None)
            render_summary(state["records"], state["ignored"])

    # ---- 3b. folder upload ----------------------------------------------
    else:
        st.caption(
            "Pick a folder from your computer. You will get a ZIP containing an "
            "**Outputs** folder with the converted files."
        )
        try:
            folder_uploads = st.file_uploader(
                "Choose a folder",
                type=["xlsb"],
                accept_multiple_files="directory",
                key="up_folder",
            )
        except TypeError:
            st.info("Folder picking needs a newer Streamlit; select the files instead.")
            folder_uploads = st.file_uploader(
                "Select the .xlsb files",
                type=["xlsb"],
                accept_multiple_files=True,
                key="up_folder_files",
            )

        if not folder_uploads:
            if (st.session_state.get("res_folder") or {}).get("kind") == "upload":
                st.session_state.pop("res_folder", None)
        else:
            st.caption(f"{len(folder_uploads)} .xlsb file(s) received · {fmt_size(sum(u.size for u in folder_uploads))}")

            if st.button("Convert folder", type="primary", key="go_folder_upload"):
                work = new_workdir("folder")
                in_dir, out_dir = work / "in", work / "Outputs"
                in_dir.mkdir()
                out_dir.mkdir()

                jobs, used, rejected = [], set(), []
                for index, item in enumerate(folder_uploads):
                    try:
                        rel = cv.safe_relative_path(item.name)
                    except ValueError as exc:
                        rejected.append(str(exc))
                        continue
                    if rel.suffix.lower() != ".xlsb" or rel.name.startswith("~$"):
                        continue
                    # Folder uploads may carry sub-paths; keep the structure.
                    dst = cv.unique_output_path(out_dir, rel, used)
                    src = in_dir / f"{index}_{rel.name}"
                    src.write_bytes(item.getbuffer())
                    jobs.append((rel.as_posix(), src, dst, dst.relative_to(out_dir).as_posix()))

                for message in rejected:
                    st.warning(message)

                if jobs:
                    records = run_jobs(jobs, engine_choice, int(timeout_seconds))
                    st.session_state["res_folder"] = {
                        "kind": "upload",
                        "records": records,
                        "zip": make_zip(records, prefix="Outputs/") if any(r["ok"] for r in records) else b"",
                        "ignored": 0,
                    }
                else:
                    st.info("No .xlsb files were found in that folder.")

            state = st.session_state.get("res_folder")
            if state and state["kind"] == "upload":
                st.divider()
                render_summary(state["records"])
                if state["zip"]:
                    st.download_button(
                        "⬇️  Download Outputs (ZIP)",
                        data=state["zip"],
                        file_name="Outputs.zip",
                        mime="application/zip",
                        type="primary",
                        key="dl_folder_zip",
                    )
                    st.caption("Unzip it next to your original files to get the Outputs folder.")
